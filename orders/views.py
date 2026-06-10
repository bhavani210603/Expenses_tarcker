import openpyxl
from django.contrib import messages
from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse
from .models import Order
from .forms import OrderForm
import csv
# from django.shortcuts import render, get_object_or_404, redirect
# from django.http import HttpResponse
# from django.contrib import messages
# from .models import Order
# from .forms import OrderForm
# import csv

# Dashboard
def dashboard(request):
    orders = Order.objects.all()

    # Chart data — profit per order (last 10)
    recent = orders.order_by('-order_date')[:10]
    chart_labels  = [o.order_id for o in recent]
    chart_profits = [float(o.profit) for o in recent]

    context = {
        'total_orders'    : orders.count(),
        'delivered'       : orders.filter(delivery_status='delivered').count(),
        'pending_payment' : orders.filter(payment_status='awaiting').count(),
        'total_earnings'  : sum(o.bank_settlement for o in orders),
        'total_profit'    : sum(o.profit for o in orders),
        'recent_orders'   : orders.order_by('-created_at')[:5],
        'chart_labels'    : chart_labels,
        'chart_profits'   : chart_profits,
    }
    return render(request, 'orders/dashboard.html', context)

# List all orders with search + filter
def order_list(request):
    orders = Order.objects.all().order_by('-created_at')

    search   = request.GET.get('search', '')
    delivery = request.GET.get('delivery', '')
    payment  = request.GET.get('payment', '')

    if search:
        orders = orders.filter(product_name__icontains=search)
    if delivery:
        orders = orders.filter(delivery_status=delivery)
    if payment:
        orders = orders.filter(payment_status=payment)

    context = {
        'orders'          : orders,
        'search'          : search,
        'delivery'        : delivery,
        'payment'         : payment,
        'delivery_choices': Order.DELIVERY_CHOICES,
        'payment_choices' : Order.PAYMENT_CHOICES,
    }
    return render(request, 'orders/order_list.html', context)

# Add new order
def order_add(request):
    if request.method == 'POST':
        form = OrderForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('order_list')
    else:
        form = OrderForm()
    return render(request, 'orders/order_form.html', {'form': form, 'title': 'Add Order'})

# Edit existing order
def order_edit(request, pk):
    order = get_object_or_404(Order, pk=pk)
    if request.method == 'POST':
        form = OrderForm(request.POST, instance=order)
        if form.is_valid():
            form.save()
            return redirect('order_list')
    else:
        form = OrderForm(instance=order)
    return render(request, 'orders/order_form.html', {'form': form, 'title': 'Edit Order'})

# Delete order
def order_delete(request, pk):
    order = get_object_or_404(Order, pk=pk)
    if request.method == 'POST':
        order.delete()
        return redirect('order_list')
    return render(request, 'orders/order_confirm_delete.html', {'order': order})

# Export to CSV
def export_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="meesho_orders.csv"'

    writer = csv.writer(response)
    writer.writerow([
        'Order ID', 'Product', 'Quantity', 'Order Date',
        'Meesho Amount', 'Product Price', 'Packing Cost',
        'Platform Fees', 'Bank Settlement', 'Total Cost',
        'Profit', 'Delivery Status', 'Payment Status', 'Condition', 'Notes'
    ])

    for o in Order.objects.all().order_by('-created_at'):
        writer.writerow([
            o.order_id, o.product_name, o.quantity, o.order_date,
            o.meesho_amount, o.product_price, o.packing_cost,
            o.platform_fees, o.bank_settlement, o.total_cost,
            o.profit, o.get_delivery_status_display(),
            o.get_payment_status_display(), o.get_product_condition_display(),
            o.notes
        ])

    return response

    # Import from Excel
    def import_excel(request):
        if request.method == 'POST' and request.FILES.get('excel_file'):
            excel_file = request.FILES['excel_file']
            wb = openpyxl.load_workbook(excel_file, data_only=True)

            imported = 0
            skipped  = 0
            errors   = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))

                if not rows:
                    continue

                # Find header row (row that contains 'Order No')
                header_row_idx = None
                for i, row in enumerate(rows):
                    if any(str(cell).strip().lower() in ['order no', 'order_no'] for cell in row if cell):
                        header_row_idx = i
                        break

                if header_row_idx is None:
                    continue

                headers = [str(h).strip().lower().replace(' ', '_') if h else '' for h in rows[header_row_idx]]

                for row in rows[header_row_idx + 1:]:
                    # Skip empty rows and total rows
                    if not any(row):
                        continue
                    data = dict(zip(headers, row))

                    order_no = str(data.get('order_no', '')).strip()
                    if not order_no or order_no.lower() in ['total', 'nan', 'none', '']:
                        continue

                    # Map Excel columns to model fields
                    try:
                        def safe_float(val, default=0):
                            try:
                                f = float(val)
                                return 0 if f < 0 and val not in [None, ''] else f
                            except (TypeError, ValueError):
                                return default

                        product_price  = safe_float(data.get('product_prise') or data.get('product_price'))
                        packing        = safe_float(data.get('packging') or data.get('packing_cost'))
                        settlement     = safe_float(data.get('bank_settlement'))
                        total_cost     = safe_float(data.get('total_cost'))
                        platform_fees  = safe_float(data.get('platform_fees'))

                        # Determine delivery status
                        status_raw = str(data.get('status') or '').strip().lower()
                        if 'rto' in status_raw:
                            delivery_status = 'returned'
                        elif 'return' in status_raw:
                            delivery_status = 'returned'
                        elif 'success' in status_raw:
                            delivery_status = 'delivered'
                        elif 'cancel' in status_raw:
                            delivery_status = 'cancelled'
                        else:
                            delivery_status = 'pending'

                        # Determine payment status
                        payment_status = 'transferred' if settlement > 0 and delivery_status == 'delivered' else 'awaiting'

                        # Handle date
                        order_date = data.get('date')
                        if not order_date:
                            from datetime import date
                            order_date = date.today()

                        # Skip if order already exists
                        if Order.objects.filter(order_id=order_no).exists():
                            skipped += 1
                            continue

                        Order.objects.create(
                            order_id        = order_no,
                            product_name    = str(data.get('product_name') or '').strip() or 'Unknown',
                            quantity        = 1,
                            order_date      = order_date,
                            meesho_amount   = total_cost,
                            product_price   = product_price,
                            packing_cost    = packing,
                            platform_fees   = platform_fees,
                            bank_settlement = settlement,
                            delivery_status = delivery_status,
                            payment_status  = payment_status,
                            product_condition = 'new',
                            notes           = f"Imported from {sheet_name}",
                        )
                        imported += 1

                    except Exception as e:
                        errors.append(f"Row {order_no}: {str(e)}")
                        continue

            msg = f"Import done! {imported} orders added, {skipped} already existed."
            if errors:
                msg += f" {len(errors)} rows had errors."
            messages.success(request, msg)
            return redirect('order_list')

        return render(request, 'orders/import_excel.html')
    
# Import from Excel
def import_excel(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        import openpyxl
        from datetime import date
        from django.contrib import messages as msg

        excel_file = request.FILES['excel_file']
        wb = openpyxl.load_workbook(excel_file, data_only=True)

        imported = 0
        skipped  = 0
        errors   = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                continue

            # Find header row
            header_row_idx = None
            for i, row in enumerate(rows):
                if any(str(cell).strip().lower() in ['order no', 'order_no'] for cell in row if cell):
                    header_row_idx = i
                    break

            if header_row_idx is None:
                continue

            headers = [str(h).strip().lower().replace(' ', '_') if h else '' for h in rows[header_row_idx]]

            for row in rows[header_row_idx + 1:]:
                if not any(row):
                    continue
                data = dict(zip(headers, row))

                order_no = str(data.get('order_no', '')).strip()
                if not order_no or order_no.lower() in ['total', 'nan', 'none', '']:
                    continue

                try:
                    def safe_float(val, default=0):
                        try:
                            return max(0, float(val))
                        except (TypeError, ValueError):
                            return default

                    product_price = safe_float(data.get('product_prise') or data.get('product_price'))
                    packing       = safe_float(data.get('packging') or data.get('packing_cost'))
                    settlement    = safe_float(data.get('bank_settlement'))
                    total_cost    = safe_float(data.get('total_cost'))
                    platform_fees = safe_float(data.get('platform_fees'))

                    status_raw = str(data.get('status') or '').strip().lower()
                    if 'rto' in status_raw:
                        delivery_status = 'returned'
                    elif 'return' in status_raw:
                        delivery_status = 'returned'
                    elif 'success' in status_raw:
                        delivery_status = 'delivered'
                    elif 'cancel' in status_raw:
                        delivery_status = 'cancelled'
                    else:
                        delivery_status = 'pending'

                    payment_status = 'transferred' if settlement > 0 and delivery_status == 'delivered' else 'awaiting'

                    order_date = data.get('date')
                    if not order_date:
                        order_date = date.today()

                    if Order.objects.filter(order_id=order_no).exists():
                        skipped += 1
                        continue

                    Order.objects.create(
                        order_id          = order_no,
                        product_name      = str(data.get('product_name') or '').strip() or 'Unknown',
                        quantity          = 1,
                        order_date        = order_date,
                        meesho_amount     = total_cost,
                        product_price     = product_price,
                        packing_cost      = packing,
                        platform_fees     = platform_fees,
                        bank_settlement   = settlement,
                        delivery_status   = delivery_status,
                        payment_status    = payment_status,
                        product_condition = 'new',
                        notes             = f"Imported from {sheet_name}",
                    )
                    imported += 1

                except Exception as e:
                    errors.append(f"Row {order_no}: {str(e)}")
                    continue

        from django.contrib import messages
        result_msg = f"Import done! {imported} orders added, {skipped} already existed."
        if errors:
            result_msg += f" {len(errors)} rows had errors."
        messages.success(request, result_msg)
        return redirect('order_list')

    return render(request, 'orders/import_excel.html')

def delete_all_orders(request):
    Order.objects.all().delete()
    messages.success(request, "All orders deleted successfully.")
    return redirect('order_list')